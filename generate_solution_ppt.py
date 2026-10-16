"""
generate_solution_ppt.py
--------------------------
Generates the complete Master Unified Solution & Technical Architecture Flowchart PPT (.pptx)
Layout:
- Top 5 Flowchart Columns: Ingestion -> Fusion -> ML Engine -> Dynamic Risk (Dijkstra) -> 3D GIS Command
- Bottom 2 Adjusted Boxes: Multi-Channel Alerts & Off-Grid Mesh | Enterprise Tech Stack
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Colors ──
NAVY = RGBColor(15, 23, 42)
BLUE_700 = RGBColor(29, 78, 216)
BLUE_100 = RGBColor(239, 246, 255)
TEAL_700 = RGBColor(15, 118, 110)
TEAL_100 = RGBColor(240, 253, 250)
ORANGE_600 = RGBColor(234, 88, 12)
ORANGE_100 = RGBColor(255, 247, 237)
PURPLE_700 = RGBColor(109, 40, 217)
PURPLE_100 = RGBColor(250, 245, 255)
INDIGO_800 = RGBColor(55, 48, 163)
INDIGO_100 = RGBColor(238, 242, 255)
WHITE = RGBColor(255, 255, 255)
SLATE_900 = RGBColor(15, 23, 42)
SLATE_600 = RGBColor(71, 85, 105)
SLATE_100 = RGBColor(241, 245, 249)

def add_rounded_rect(slide, left, top, width, height, fill_color, border_color=None, border_width=Pt(1.2)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    return shape

def add_arrow(slide, left, top, width, height, color=RGBColor(100, 116, 139)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_txt(slide, left, top, width, height, text, font_size=10, bold=False, color=SLATE_900, align=PP_ALIGN.LEFT):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Arial"
    p.alignment = align
    return tx

def build_master_ppt(out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Header
    add_rounded_rect(slide, Inches(0.4), Inches(0.2), Inches(12.533), Inches(0.65), NAVY)
    add_txt(slide, Inches(0.6), Inches(0.24), Inches(3.2), Inches(0.55), "Team 404: The Optimistics", 13, True, WHITE)
    add_txt(slide, Inches(3.8), Inches(0.22), Inches(5.8), Inches(0.55), "END-TO-END SOLUTION & TECHNICAL ARCHITECTURE FLOW", 16, True, WHITE, PP_ALIGN.CENTER)
    add_txt(slide, Inches(9.8), Inches(0.24), Inches(2.9), Inches(0.55), "Smart India Hackathon 2026", 12, True, RGBColor(251, 146, 60), PP_ALIGN.RIGHT)

    # 5 Flowchart Columns
    col_w = Inches(2.25)
    gap = Inches(0.25)
    arrow_w = Inches(0.18)
    start_x = Inches(0.4)
    col_top = Inches(0.95)
    col_h = Inches(4.5)

    cols = [
        ("① DATA INGESTION", BLUE_700, BLUE_100, [
            "• Satellite Data\n  (SMAP, SRTM, InSAR)",
            "• Weather Feeds\n  (IMD, Open-Meteo)",
            "• Terrain & Rivers\n  (CWC Gauges, IoT)",
            "• Citizen Reports\n  (Geotagged Photos)"
        ]),
        ("② DATA FUSION", TEAL_700, TEAL_100, [
            "• Data Cleaning &\n  30m grid alignment",
            "• Multi-Source\n  Spatial Overlay",
            "• Feature Engineering\n  (9 Core Risk Factors)",
            "  - Rain 1d/3d/7d\n  - Soil Moisture\n  - Slope & InSAR Slip"
        ]),
        ("③ AI/ML ENGINE", ORANGE_600, ORANGE_100, [
            "• ML Ensemble Stack\n  (RF, XGBoost, NN)",
            "• Probabilistic Risk\n  P(risk) = 0.82",
            "• 4-Level Severity:\n  High | Mod | Low | Min",
            "• Verified Benchmark\n  ROC-AUC = 0.894"
        ]),
        ("④ DYNAMIC RISK", PURPLE_700, PURPLE_100, [
            "• Dynamic Cost Surface:\n  Cost = Dist*(1+R*12)",
            "• 3D Dijkstra Graph\n  Safe Path Router",
            "• 400m Hazard Buffer\n  Auto-Bypass",
            "• Live Dynamic Reroute\n  (Sub-second response)"
        ]),
        ("⑤ 3D COMMAND", INDIGO_800, INDIGO_100, [
            "• 3D WebGL Digital Twin\n  (MapLibre / Deck.gl)",
            "• Tactical Hazards Layer\n  (Roads, Bridges, Slips)",
            "• Safe Shelter Registry\n  & Capacity Tracking",
            "• SDRF / NDRF Dispatch\n  Coordination Hub"
        ])
    ]

    for i, (col_title, header_col, bg_col, bullets) in enumerate(cols):
        cx = start_x + i * (col_w + gap)

        # Header Box
        add_rounded_rect(slide, cx, col_top, col_w, Inches(0.45), header_col)
        add_txt(slide, cx, col_top + Inches(0.08), col_w, Inches(0.35), col_title, 11, True, WHITE, PP_ALIGN.CENTER)

        # Body Box
        add_rounded_rect(slide, cx, col_top + Inches(0.48), col_w, col_h - Inches(0.48), bg_col, header_col)

        # Bullet Text
        by = col_top + Inches(0.6)
        for bullet in bullets:
            add_txt(slide, cx + Inches(0.1), by, col_w - Inches(0.2), Inches(0.7), bullet, 8.5, True, SLATE_900)
            by += Inches(0.85)

        # Draw connecting arrow to next column
        if i < 4:
            ax = cx + col_w + Inches(0.035)
            ay = col_top + Inches(2.1)
            add_arrow(slide, ax, ay, arrow_w, Inches(0.25), header_col)

    # ── Bottom Section (2 Adjusted Boxes) ──
    bot_top = Inches(5.6)
    bot_h = Inches(1.55)

    # Left: Multi-Channel Alerts & Off-Grid Mesh
    left_w = Inches(7.5)
    add_rounded_rect(slide, Inches(0.4), bot_top, left_w, bot_h, SLATE_100, RGBColor(225, 29, 72))
    add_txt(slide, Inches(0.6), bot_top + Inches(0.08), left_w - Inches(0.4), Inches(0.3),
            "MULTI-CHANNEL ALERTS & OFF-GRID MESH CONNECTIVITY", 11, True, RGBColor(190, 18, 60))

    alert_items = [
        "• Citizen Mobile App Push Notifications (Evacuation alerts, SOS emergency trigger)",
        "• Regional Multilingual SMS & IVR Broadcast (Assamese, Bengali, Hindi, etc.)",
        "• High-Decibel Audio Sirens at high-risk landslide & flood choke points",
        "• Offline BLE / LoRa Mesh Network: P2P relay when telecom towers & power fail"
    ]
    ay = bot_top + Inches(0.4)
    for a in alert_items:
        add_txt(slide, Inches(0.6), ay, left_w - Inches(0.4), Inches(0.25), a, 8.5, False, SLATE_900)
        ay += Inches(0.26)

    # Right: Production Tech Stack
    right_x = Inches(8.1)
    right_w = Inches(4.833)
    add_rounded_rect(slide, right_x, bot_top, right_w, bot_h, SLATE_100, BLUE_700)
    add_txt(slide, right_x + Inches(0.2), bot_top + Inches(0.08), right_w - Inches(0.4), Inches(0.3),
            "ENTERPRISE TECH STACK", 11, True, BLUE_700, PP_ALIGN.CENTER)

    stack_rows = [
        "Python 3.14 (FastAPI, Scikit-learn, GDAL)  |  Spring Boot 3.3",
        "Next.js 14 / TypeScript (MapLibre 3D WebGL, Deck.gl)",
        "PostgreSQL + PostGIS (Spatial Indexing)  |  Supabase Storage",
        "NVIDIA NIM AI (ResQBot Voice & Chat)  |  BLE Mesh / LoRa"
    ]
    sy = bot_top + Inches(0.4)
    for s in stack_rows:
        add_txt(slide, right_x + Inches(0.2), sy, right_w - Inches(0.4), Inches(0.25), s, 8, True, SLATE_900, PP_ALIGN.CENTER)
        sy += Inches(0.26)

    prs.save(out_path)
    print("Master PPT successfully saved at:", out_path)

if __name__ == "__main__":
    build_master_ppt(r"c:\Users\KIIT\Desktop\resQVerse\NER-Sentinel_Master_Architecture.pptx")
